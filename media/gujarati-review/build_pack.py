import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

here = sys.argv[1]; out = sys.argv[2]
rows = json.load(open(f"{here}/rows.json"))
by = {r["key"]: r for r in rows}

SECTIONS = [
  ("Menu and footer (every page)", ["shell.", "page."]),
  ("Home page", ["home."]),
  ("What's on page", ["whatsOn.", "heading.", "day."]),
  ("Mandir list and map", ["mandirs.", "region.", "map."]),
  ("A mandir's page", ["mandir."]),
  ("An event's page", ["event."]),
  ("Festival dates and festival pages", ["festivals.", "festival."]),
  ("Season pages (Diwali week, Navratri, Holi…)", ["season."]),
  ("Seva sign-up (volunteers)", ["seva.", "slot."]),
  ("'Tell us about an event' form", ["report."]),
  ("Small pieces used on many pages", ["kind.", "source.", "time.", "repeat.", "calendar.", "distance.", "share.", "subscribe.", "countdown.", "load."]),
]
def section(key):
  for name, prefixes in SECTIONS:
    if any(key.startswith(p) for p in prefixes): return name
  return "Other"

# Example values for {placeholders}, so each line can be read as it shows.
FEST = {"name": "દિવાળી", "year": "2026", "dates": "રવિ 8 નવે 2026"}
EX = {"name": "Shree Ram Mandir", "town": "Leicester", "count": "3", "total": "106", "shown": "12",
      "miles": "10", "place": "LE4", "region": "ઈસ્ટ મિડલેન્ડ્સ", "type": "તહેવાર", "query": "Lester",
      "title": "Diwali celebrations", "when": "રવિ 8 નવે, સાંજે 7", "date": "રવિ 18 ઑક્ટો 2026",
      "day": "રવિવારે", "ordinal": "પહેલા", "rule": "દર રવિવારે", "time": "7:30", "days": "5",
      "taken": "4", "needed": "10", "done": "4", "free": "6", "role": "Kitchen", "label": "3 મંદિરો",
      "what": "કાર્યક્રમોની યાદી", "text": "Shree Ram Mandir, Leicester: કાર્યક્રમો", "period": "દિવાળી સપ્તાહ",
      "dates": "રવિ 8 નવે 2026"}
EX_EN = dict(EX, region="East Midlands", type="Festival", when="Sun 8 Nov, 7pm",
  date="Sun 18 Oct 2026", day="Sunday", ordinal="First", rule="Every Sunday", label="3 mandirs",
  what="event list", text="Shree Ram Mandir, Leicester: what's on", period="Diwali week",
  dates="Sun 8 Nov 2026")
FEST_EN = dict(FEST, name="Diwali", dates="Sun 8 Nov 2026")
def filled(key, text, lang="gu"):
  vals = dict(EX if lang == "gu" else EX_EN)
  if key.startswith(("festival.", "season.")): vals.update(FEST if lang == "gu" else FEST_EN)
  return re.sub(r"\{(\w+)\}", lambda m: vals.get(m.group(1), m.group(0)), text)

# Pieces of one sentence split around a link: show the whole sentence, as
# the page joins them. A tuple is a literal (English, Gujarati).
SP = (" ", " ")
GROUPS = [
  (["whatsOn.emptyBefore", "whatsOn.emptyLink", "whatsOn.emptyAfter"], "\"mandir list\" is a link"),
  (["mandir.claimedBefore", SP, "mandir.logIn", SP, "mandir.claimedAfter"], "\"Log in\" is a link"),
  (["mandir.untilThenCheck", SP, ("leedsmandir.org.uk", "leedsmandir.org.uk"), SP, "mandir.orNoticeboard"], "the web address is a link to the mandir's website"),
  (["seva.takenOf", SP, "seva.placesTaken"], "on an event's page; the numbers are in bold"),
  (["seva.takenOf", SP, "seva.pickBelow"], "on the seva page; the numbers are in bold"),
  (["festival.appearBefore", SP, "festival.whatsOnLink", SP, "festival.orFind", SP, "festival.nearestMandir", "festival.appearAfter"], "on a festival page with no events yet; two links"),
  (["season.notListed", SP, "season.festivalDates", "festival.appearAfter"], "\"festival dates\" is a link"),
  (["season.adding", SP, "festival.nearestMandir", "festival.appearAfter", SP, "season.committee", SP, "season.addProgramme"], "on a season page with no events yet; two links"),
  (["whatsOn.count.other", "whatsOn.countNear", "whatsOn.countRegion", "whatsOn.countType", "whatsOn.nearestFirst"], "the count at the top of What's on, with every filter used; only the parts that apply are shown"),
]
reads = {}
for parts, note in GROUPS:
  def text(lang):
    out = []
    for p in parts:
      if isinstance(p, tuple): out.append(p[0] if lang == "en" else p[1])
      else: out.append(filled(p, by[p][lang], lang))
    return re.sub(" +", " ", "".join(out)).strip()
  label = f"Whole sentence ({note})\nEN: {text('en')}\nગુ: {text('gu')}"
  for p in parts:
    if isinstance(p, str) and label not in reads.get(p, []):
      reads.setdefault(p, []).append(label)

NOTES = {
  "event.repeats": "\"દર વખતે\" means \"every time\". Is that right for a label saying an event repeats (weekly etc.)?",
  "festival.summary.diwali": "\"દીવાઓનો તહેવાર\" or \"પ્રકાશનો તહેવાર\"?",
  "home.step1": "Here \"વોટ્સએપ\", but the 'Tell us' form keeps \"WhatsApp\". Which should we use everywhere?",
  "home.step3": "Same question: \"વોટ્સએપ\" or \"WhatsApp\"?",
  "report.intro": "Same question: \"WhatsApp\" or \"વોટ્સએપ\"?",
  "report.whatExample": "The committee's event form says \"દા.ત. દિવાળી ઉજવણી\" (see the Other bits sheet). Which is better?",
  "season.calendarTitle": "Spelled \"કૅલેન્ડર\" here but \"કેલેન્ડર\" elsewhere. Which spelling?",
  "season.calendarBody": "\"કૅલેન્ડર\" vs \"કેલેન્ડર\", and \"આપોઆપ\" here vs \"આપમેળે\" elsewhere. Which?",
  "mandir.calendarBody": "\"આપમેળે\" here, \"આપોઆપ\" on season pages. Which?",
  "festivals.calendarBody": "\"આપમેળે\" here, \"આપોઆપ\" on season pages. Which?",
  "whatsOn.description": "\"હિન્દુ\" here, \"હિંદુ\" on season pages. Which spelling? Also \"યુકેભરના\" vs \"યુકેભરનાં\".",
  "mandirs.description": "\"હિન્દુ\" or \"હિંદુ\"?",
  "season.diwali.title": "\"યુકેનાં મંદિરોમાં\" here; other pages say \"યુકેભરના મંદિરો\" (without the anusvara). Which agreement is right?",
  "time.night": "Used for times after 8pm, e.g. \"રાત્રે 9\". And \"સાંજે\" for 4pm to 8pm. Do those boundaries feel right?",
  "time.morning": "Used for 4am to noon, e.g. \"સવારે 7:30\".",
  "time.afternoon": "Used for noon to 4pm, e.g. \"બપોરે 2\".",
  "mandir.tradition": "Shows e.g. \"Swaminarayan (BAPS)\". Is \"સંપ્રદાય\" the right label?",
  "calendar.gujarati": "\"પંચાંગ\" for \"calendar\" here (the tradition a mandir follows). Right word?",
  "shell.whatsOn": "\"What's on\" is the site's main section. \"કાર્યક્રમો\" (programmes): natural?",
  "home.tagline": "The site's tagline, shown large on the home page and in shared links.",
  "map.missing.other": "Plural: \"તે યાદીમાં છે\". Better \"તે બધાં યાદીમાં છે\"?",
  "slot.youreOn": "Shown on a seva slot the volunteer took.",
  "festival.summary.bhai-dooj": "Is Bhai Bij two days after Diwali in the Gujarati calendar (this year Wed 11 Nov? or Tue 10?)",
  "countdown.days": "Means \"5 days to go\".",
  "whatsOn.countType": "Today the type in brackets shows in English on Gujarati pages; we'll fix that.",
  "festival.summary.uttarayan": "English says khichdi; Gujarati says \"ખીચડો\" (the Uttarayan dish). OK?",
}
# Plurals and placeholders: say which words are filled in.
for r in rows:
  ph = sorted(set(re.findall(r"\{(\w+)\}", r["en"])))
  r["placeholders"] = ph

wb = Workbook()
ARIAL = "Arial"
base = Font(name=ARIAL, size=12)
bold = Font(name=ARIAL, size=12, bold=True)
head = Font(name=ARIAL, size=12, bold=True, color="FFFFFF")
headfill = PatternFill("solid", fgColor="0A3B3A")
fillin = PatternFill("solid", fgColor="FFF4CC")
secfill = PatternFill("solid", fgColor="E6F0EE")
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D9D4C7")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---- Start here ----
ws = wb.active; ws.title = "Start here"
lines = [
  ("Toran: checking the Gujarati", bold),
  ("", None),
  ("Toran is a website that shows what's on at mandirs across the UK. It has a Gujarati version that isn't switched on yet, because a native speaker hasn't checked it. Thank you for helping.", None),
  ("", None),
  ("What to do", bold),
  ("1. Open the \"Lines to check\" sheet. Each row has the English, our draft Gujarati, and how the line reads with real words filled in.", None),
  ("2. In the yellow \"OK?\" column, choose OK, Change or Not sure.", None),
  ("3. If you choose Change, write your better Gujarati in \"Your Gujarati\". Anything else goes in \"Comment\".", None),
  ("4. Some rows have a question from us in \"Things to check\". Please answer it in \"Comment\".", None),
  ("5. The \"Other bits\" sheet has dates, times and a few other pieces. Please check those too.", None),
  ("6. Please answer the general questions below in the yellow boxes, then save the file and send it back.", None),
  ("", None),
  ("Please keep", bold),
  ("• Words in {curly brackets} exactly as they are, for example {name} or {count}. The site puts the real word there: a mandir's name, a number, a date. You can move them within the sentence.", None),
  ("• \"Toran\" in English letters: it's the site's name.", None),
  ("• The meaning, but don't feel bound to the English wording. Natural, warm Gujarati that a grandparent would be comfortable reading is the goal.", None),
  ("", None),
  ("Example of a filled-in row", bold),
  ("OK? = Change  ·  Your Gujarati = \"દર અઠવાડિયે\"  ·  Comment = \"Sounds more natural\"", None),
  ("", None),
  ("General questions", bold),
]
for i, (t, f) in enumerate(lines, 1):
  c = ws.cell(row=i, column=1, value=t); c.font = f or base; c.alignment = wrap
r0 = len(lines) + 1
QUESTIONS = [
  "The draft says \"તમે\" (you) throughout. Would \"આપ\" be better for a mandir audience, or is \"તમે\" right for a website?",
  "Some English words are kept in English letters (Google, Apple, Outlook, WhatsApp, DBS, postcode examples). Is that OK?",
  "Are the event types right and natural: આરતી કે દર્શન, ભજન કે કીર્તન, કથા કે પ્રવચન, વર્ગ, યુવાનો, સેવા દિવસ, સમુદાય?",
  "Anything that sounds like a machine translation, too formal, or too casual?",
]
for j, q in enumerate(QUESTIONS):
  ws.cell(row=r0 + 2*j, column=1, value=f"{j+1}. {q}").font = base
  ws.cell(row=r0 + 2*j, column=1).alignment = wrap
  a = ws.cell(row=r0 + 2*j + 1, column=1, value=""); a.fill = fillin; a.border = box; a.alignment = wrap
  ws.row_dimensions[r0 + 2*j + 1].height = 48
ws.column_dimensions["A"].width = 110

# ---- Lines to check ----
ls = wb.create_sheet("Lines to check")
cols = [("#", 6), ("Where it shows", 22), ("English", 42), ("Draft Gujarati", 42), ("How it reads on the page", 46), ("Things to check", 34), ("OK?", 11), ("Your Gujarati", 40), ("Comment", 32), ("key", 10)]
for j, (name, w) in enumerate(cols, 1):
  c = ls.cell(row=1, column=j, value=name); c.font = head; c.fill = headfill; c.alignment = wrap; c.border = box
  ls.column_dimensions[c.column_letter].width = w
ls.freeze_panes = "C2"
order = {name: i for i, (name, _) in enumerate(SECTIONS)}
ordered = sorted(enumerate(rows), key=lambda t: (order.get(section(t[1]["key"]), 99), t[0]))
dv = DataValidation(type="list", formula1='"OK,Change,Not sure"', allow_blank=True)
ls.add_data_validation(dv)
row = 2; n = 0; last = None
for _, r in ordered:
  sec = section(r["key"])
  if sec != last:
    c = ls.cell(row=row, column=1, value=sec); c.font = bold
    for j in range(1, len(cols)+1): ls.cell(row=row, column=j).fill = secfill
    ls.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(cols)-1)
    row += 1; last = sec
  n += 1
  how = []
  if r["placeholders"]:
    how.append(f"EN: {filled(r['key'], r['en'], 'en').strip()}\nગુ: {filled(r['key'], r['gu']).strip()}")
  how += reads.get(r["key"], [])
  vals = [n, sec.split(" (")[0], r["en"], r["gu"], "\n\n".join(how), NOTES.get(r["key"], ""), "", "", "", r["key"]]
  for j, v in enumerate(vals, 1):
    c = ls.cell(row=row, column=j, value=v); c.font = base; c.alignment = wrap; c.border = box
    if j in (7, 8, 9): c.fill = fillin
  dv.add(ls.cell(row=row, column=7))
  row += 1
ls.column_dimensions["J"].hidden = True
ls.auto_filter.ref = f"A1:I{row-1}"

# ---- Other bits ----
ob = wb.create_sheet("Other bits")
ocols = [("#", 6), ("What it is", 40), ("English", 30), ("Draft Gujarati", 34), ("OK?", 11), ("Your Gujarati", 34), ("Comment", 32)]
for j, (name, w) in enumerate(ocols, 1):
  c = ob.cell(row=1, column=j, value=name); c.font = head; c.fill = headfill; c.alignment = wrap; c.border = box
  ob.column_dimensions[c.column_letter].width = w
ob.freeze_panes = "B2"
dates = json.load(open(f"{here}/dates.json"))
other = []
import datetime
for d in dates:
  day = datetime.date.fromisoformat(d["iso"])
  en_short = day.strftime("%a %-d %b")
  other.append(("A date, short (in lists and on date tiles)", en_short, d["short"]))
  other.append(("A date with the year", f"{en_short} {day.year}", d["withYear"]))
other.append(("A month heading", "November 2026", dates[0]["month"]))
other.append(("A month heading", "October 2026", dates[2]["month"]))
other.append(("Date and time together (event lists)", "Sun 8 Nov, 7:30pm", "રવિ 8 નવે, સાંજે 7:30"))
other.append(("Date and time, morning", "Sat 17 Oct, 10am", "શનિ 17 ઑક્ટો, સવારે 10"))
other.append(("Committee's event form: example for the Gujarati name", "e.g. Diwali celebrations", "દા.ત. દિવાળી ઉજવણી"))
other.append(("The button to switch language (on English pages)", "Gujarati", "ગુજરાતી"))
for i, (what, en, gu) in enumerate(other, 1):
  for j, v in enumerate([i, what, en, gu, "", "", ""], 1):
    c = ob.cell(row=i+1, column=j, value=v); c.font = base; c.alignment = wrap; c.border = box
    if j in (5, 6, 7): c.fill = fillin
dv2 = DataValidation(type="list", formula1='"OK,Change,Not sure"', allow_blank=True)
ob.add_data_validation(dv2); dv2.add(f"E2:E{len(other)+1}")
n2 = len(other)+3
ob.cell(row=n2, column=2, value="Dates and month names come from the phone or computer's own Gujarati settings, so we can only change them if something reads wrongly.").font = Font(name=ARIAL, size=11, italic=True)
ob.cell(row=n2, column=2).alignment = wrap
ob.merge_cells(start_row=n2, start_column=2, end_row=n2, end_column=7)
ob.row_dimensions[n2].height = 32

wb.save(out)
print(n, "lines;", len(other), "other bits")
